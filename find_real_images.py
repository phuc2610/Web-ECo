import docx

doc = docx.Document('Website-ban-quan-ao.docx')

print('Finding actual image/drawing elements in paragraphs...')
actual_img_paragraphs = []

for i, p in enumerate(doc.paragraphs):
    # Check for <w:drawing> or <w:pict>
    drawings = p._element.xpath('.//w:drawing | .//w:pict')
    if drawings:
        caption = ''
        if i + 1 < len(doc.paragraphs):
            caption = doc.paragraphs[i + 1].text.strip()
        print(f'Paragraph {i}: text="{p.text.strip()[:40]}", next_p="{caption}"')
        actual_img_paragraphs.append((i, p.text.strip(), caption))

print(f'Total actual image paragraphs found: {len(actual_img_paragraphs)}')
