import docx
import os

doc = docx.Document('Website-ban-quan-ao.docx')

with open('docx_analysis.txt', 'w', encoding='utf-8') as f:
    f.write(f'Total paragraphs: {len(doc.paragraphs)}\n')
    f.write(f'Total tables: {len(doc.tables)}\n')
    f.write(f'Total sections: {len(doc.sections)}\n\n')
    
    f.write('=== ALL PARAGRAPHS ===\n')
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        has_image = 'graphic' in p._p.xml or 'drawing' in p._p.xml
        img_info = ' [HAS_IMAGE]' if has_image else ''
        if txt or has_image:
            f.write(f'[{i}] ({p.style.name}){img_info}: {txt}\n')
            
    f.write('\n=== ALL TABLES ===\n')
    for ti, t in enumerate(doc.tables):
        f.write(f'\n--- Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) ---\n')
        for ri, row in enumerate(t.rows):
            row_txt = [c.text.replace('\n', ' ').strip() for c in row.cells]
            f.write(f'  Row {ri}: ' + ' | '.join(row_txt) + '\n')

print('Wrote docx_analysis.txt')
