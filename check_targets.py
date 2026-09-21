import docx

doc = docx.Document('Website-ban-quan-ao.docx')

def show_range(start, end, label):
    print(f'=== {label} (P{start} - P{end}) ===')
    for i in range(start, min(end, len(doc.paragraphs))):
        txt = doc.paragraphs[i].text.strip()
        if txt:
            print(f'[{i}] ({doc.paragraphs[i].style.name}): {txt[:80]}')

show_range(845, 875, 'Chapter 3 - Section 2.3 & 2.4')
show_range(1320, 1340, 'Chapter 5 - UI Screenshots')
show_range(1390, 1420, 'Chapter 5 - Admin section')
