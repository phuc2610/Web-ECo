import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_run_font(run, font_name="Times New Roman", size_pt=13, bold=False, italic=False):
    run.font.name = font_name
    if size_pt:
        run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

doc = docx.Document('Website-ban-quan-ao.docx')

# Paragraph 59: Heading
p59 = doc.paragraphs[59]
p59.text = ""
p59.alignment = WD_ALIGN_PARAGRAPH.CENTER
p59.paragraph_format.space_before = Pt(12)
p59.paragraph_format.space_after = Pt(12)
r59 = p59.add_run("LỜI NÓI ĐẦU VÀ LỜI CẢM ƠN")
set_run_font(r59, font_name="Times New Roman", size_pt=14, bold=True)

# Paragraph 60: blank separator
p60 = doc.paragraphs[60]
p60.text = ""

# Paragraph 61: Đoạn 1 - Bối cảnh & Xu hướng
p61 = doc.paragraphs[61]
p61.text = ""
p61.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p61.paragraph_format.line_spacing = 1.3
p61.paragraph_format.space_after = Pt(6)
r61 = p61.add_run(
    "Trong kỷ nguyên số hóa và sự phát triển vượt bậc của cuộc Cách mạng Công nghiệp 4.0, "
    "thương mại điện tử đã trở thành một phần thiết yếu của nền kinh tế toàn cầu cũng như đời sống xã hội. "
    "Thói quen tiêu dùng, đặc biệt là thế hệ trẻ yêu công nghệ, đã có bước chuyển dịch mạnh mẽ từ mua sắm "
    "tại các cửa hàng truyền thống sang các nền tảng số trực tuyến tiện lợi và nhanh chóng. Đối với ngành "
    "bán lẻ thiết bị công nghệ – bao gồm máy tính xách tay (Laptop), điện thoại thông minh (Smartphone), "
    "máy tính bảng và linh kiện phụ kiện số – xu hướng này càng thể hiện rõ nét khi nhu cầu tra cứu cấu hình "
    "chi tiết, so sánh giá cả theo thời gian thực và đặt mua trực tuyến ngày càng tăng cao."
)
set_run_font(r61, font_name="Times New Roman", size_pt=13)

# Paragraph 62: blank
p62 = doc.paragraphs[62]
p62.text = ""

# Paragraph 63: Đoạn 2 - Thực trạng & Thách thức
p63 = doc.paragraphs[63]
p63.text = ""
p63.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p63.paragraph_format.line_spacing = 1.3
p63.paragraph_format.space_after = Pt(6)
r63 = p63.add_run(
    "Tuy nhiên, nhiều cửa hàng kinh doanh thiết bị công nghệ quy mô vừa và nhỏ hiện nay vẫn đang vận hành "
    "theo phương thức bán lẻ thủ công hoặc phụ thuộc hoàn toàn vào các trang mạng xã hội. Phương thức này bộc lộ "
    "rất nhiều hạn chế: khó khăn trong việc phân loại và quản lý danh mục sản phẩm phức tạp với nhiều biến thể phần "
    "cứng (CPU, dung lượng RAM, ổ cứng ROM, phiên bản màu sắc), không thể đồng bộ chính xác số lượng tồn kho theo "
    "thời gian thực, dễ xảy ra thất thoát, nhầm lẫn đơn đặt hàng và đặc biệt là thiếu cơ chế tự động hóa quy trình "
    "thanh toán trực tuyến an toàn, tin cậy."
)
set_run_font(r63, font_name="Times New Roman", size_pt=13)

# Paragraph 64: blank
p64 = doc.paragraphs[64]
p64.text = ""

# Paragraph 65: Đoạn 3 - Giải pháp đề tài Minh Tuấn Shop
p65 = doc.paragraphs[65]
p65.text = ""
p65.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p65.paragraph_format.line_spacing = 1.3
p65.paragraph_format.space_after = Pt(6)
r65 = p65.add_run(
    "Xuất phát từ thực tế đó, đề tài “Xây dựng website thương mại điện tử kinh doanh máy tính và điện thoại – "
    "Minh Tuấn Shop” được nhóm tác giả nghiên cứu và phát triển nhằm mang đến giải pháp chuyển đổi số bán hàng "
    "toàn diện. Đối với khách hàng, hệ thống mang đến giao diện trực quan, hỗ trợ tìm kiếm sản phẩm thông minh, "
    "bộ lọc đa năng theo thương hiệu và cấu hình, tùy chọn thông số kỹ thuật chi tiết, quản lý giỏ hàng, danh sách "
    "yêu thích và đặc biệt là giải pháp thanh toán quét mã VietQR tự động 24/7 (tích hợp cổng SePay) bên cạnh cổng "
    "thanh toán quốc tế Stripe và COD. Đối với quản trị viên (Admin Portal), hệ thống cung cấp công cụ trực quan để "
    "quản lý danh mục sản phẩm, biến thể cấu hình, cập nhật bảng giá, kiểm soát trạng thái đơn hàng và theo dõi "
    "doanh thu bán hàng tức thời."
)
set_run_font(r65, font_name="Times New Roman", size_pt=13)

# Paragraph 66: blank
p66 = doc.paragraphs[66]
p66.text = ""

# Paragraph 67: Đoạn 4 - Công nghệ MERN Stack & Dịch vụ tích hợp
p67 = doc.paragraphs[67]
p67.text = ""
p67.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p67.paragraph_format.line_spacing = 1.3
p67.paragraph_format.space_after = Pt(6)
r67 = p67.add_run(
    "Hệ thống được thiết kế theo kiến trúc 3 lớp hiện đại, phát triển trên nền tảng ngăn xếp công nghệ MERN Stack "
    "(MongoDB, ExpressJS, ReactJS, NodeJS) kết hợp với Tailwind CSS, dịch vụ đám mây Cloudinary (lưu trữ và phân phối "
    "hình ảnh đa phương tiện tốc độ cao), cơ chế Webhook SePay (tự động nhận diện biến động số dư ngân hàng Việt Nam "
    "để kích hoạt đơn hàng trong vài giây) và Stripe API. Nền tảng công nghệ này đảm bảo hệ thống vận hành mượt mà, "
    "hiệu năng tối ưu, tính bảo mật cao và sẵn sàng mở rộng quy mô trong tương lai."
)
set_run_font(r67, font_name="Times New Roman", size_pt=13)

# Paragraph 68: blank
p68 = doc.paragraphs[68]
p68.text = ""

# Paragraph 69: Đoạn 5 - Bố cục đồ án
p69 = doc.paragraphs[69]
p69.text = ""
p69.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p69.paragraph_format.line_spacing = 1.3
p69.paragraph_format.space_after = Pt(6)
r69 = p69.add_run(
    "Cuốn báo cáo tốt nghiệp này trình bày một cách có hệ thống toàn bộ quy trình xây dựng dự án: từ khảo sát nghiệp vụ, "
    "phân tích yêu cầu chức năng – phi chức năng, thiết kế kiến trúc và mô hình dữ liệu (ERD, Use Case), đến cài đặt "
    "hiện thực hóa các chức năng thực tế và đánh giá kết quả kiểm thử hệ thống."
)
set_run_font(r69, font_name="Times New Roman", size_pt=13)

# Paragraph 70: blank
p70 = doc.paragraphs[70]
p70.text = ""

# Paragraph 71: Lời tri ân GVHD và Nhà trường
p71 = doc.paragraphs[71]
p71.text = ""
p71.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p71.paragraph_format.line_spacing = 1.3
p71.paragraph_format.space_after = Pt(12)
r71 = p71.add_run(
    "Chúng em xin bày tỏ lòng biết ơn sâu sắc và chân thành nhất tới Thầy Bùi Xuân Cảnh – Giảng viên hướng dẫn, "
    "người đã luôn tận tình chỉ bảo, định hướng và đóng góp những nhận xét chuyên môn quý báu giúp chúng em hoàn "
    "thiện đề tài một cách tốt nhất. Chúng em cũng xin gửi lời tri ân đến toàn thể quý Thầy/Cô Khoa Công nghệ Thông tin "
    "– Trường Đại học Lạc Hồng đã tận tâm giảng dạy, trang bị những tri thức vững chắc cho chúng em trong suốt khóa học."
)
set_run_font(r71, font_name="Times New Roman", size_pt=13)

# Paragraph 72: Ký tên sinh viên thực hiện (Căn phải)
p72 = doc.paragraphs[72]
p72.text = ""
p72.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p72.paragraph_format.space_after = Pt(4)
r72 = p72.add_run("Sinh viên thực hiện:")
set_run_font(r72, font_name="Times New Roman", size_pt=13, bold=True, italic=True)

# Paragraph 73: Tên 2 sinh viên (Căn phải)
p73 = doc.paragraphs[73]
p73.text = ""
p73.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p73.paragraph_format.space_after = Pt(18)
r73_1 = p73.add_run("Trần Tuần Lương (122000776)\nNguyễn Văn Đại (122001433)")
set_run_font(r73_1, font_name="Times New Roman", size_pt=13, bold=True)

doc.save('Website-ban-quan-ao.docx')
print("Successfully updated LỜI NÓI ĐẦU VÀ LỜI CẢM ƠN in Website-ban-quan-ao.docx!")
