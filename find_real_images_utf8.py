import docx

doc = docx.Document('Website-ban-quan-ao.docx')

with open('real_images.txt', 'w', encoding='utf-8') as out:
    for i, p in enumerate(doc.paragraphs):
        drawings = p._element.xpath('.//w:drawing | .//w:pict')
        if drawings:
            # find surrounding text
            surrounding = []
            for delta in [-2, -1, 0, 1, 2]:
                idx = i + delta
                if 0 <= idx < len(doc.paragraphs):
                    txt = doc.paragraphs[idx].text.strip()
                    if txt:
                        surrounding.append(f'[{idx}]: {txt}')
            out.write(f'--- Image at Paragraph {i} (drawings: {len(drawings)}) ---\n')
            out.write('\n'.join(surrounding) + '\n\n')

print('Wrote real_images.txt')
