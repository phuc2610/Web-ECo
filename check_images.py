import docx
import zipfile
import xml.etree.ElementTree as ET
import os

doc = docx.Document('Website-ban-quan-ao.docx')

# Check where images are placed and their following paragraphs (captions)
with open('images_info.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if 'drawing' in p._p.xml or 'graphic' in p._p.xml:
            # Found image paragraph
            f.write(f'Paragraph {i} contains an image.\n')
            # check previous 2 and next 2 paragraphs
            start = max(0, i - 2)
            end = min(len(doc.paragraphs), i + 3)
            for j in range(start, end):
                f.write(f'  [{j}] {doc.paragraphs[j].text.strip()}\n')
            f.write('\n')

print('Wrote images_info.txt')
